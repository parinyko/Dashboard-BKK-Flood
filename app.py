from flask import Flask, render_template, request, jsonify, send_file, Response
import sqlite3
import os
import time
import json
import tempfile
import re
from io import BytesIO
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

try:
    from shapely.geometry import Point, shape
    from shapely.strtree import STRtree
except Exception:
    Point = shape = STRtree = None

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "data", "sites.db")
FLOODBOARD_BASE = "https://www.floodboard.org"
FLOODBOARD_ROADS = f"{FLOODBOARD_BASE}/api/export/roads.geojson"
FLOODBOARD_REPORTS = f"{FLOODBOARD_BASE}/api/export/reports.csv"
FLOODBOARD_STATS = f"{FLOODBOARD_BASE}/api/stats"
FLOODBOARD_FEED = f"{FLOODBOARD_BASE}/api/feed"
ALLOWED_PROVINCES = ["กรุงเทพมหานคร", "ปทุมธานี", "นนทบุรี", "สมุทรปราการ"]
CACHE_DIR = os.path.join(BASE_DIR, "data", "cache")
JOB_CACHE = os.path.join(BASE_DIR, "data", "job_monitor.json")
os.makedirs(CACHE_DIR, exist_ok=True)

app = Flask(__name__)
app.config["JSON_AS_ASCII"] = False

_flood_index = None
_flood_index_loaded_at = 0
_flood_features = []


def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def row_to_dict(row):
    return dict(row)


def clean(v):
    return "" if v is None else str(v).strip()


def normalize_depth(props):
    """Return best-known water depth in cm, or None."""
    if not isinstance(props, dict):
        return None
    keys = [
        "depth", "depth_cm", "water_depth", "max_depth", "depthCm",
        "waterDepth", "maxDepth", "flood_depth", "floodDepth", "cm"
    ]
    for key in keys:
        value = props.get(key)
        try:
            if value is not None and str(value).strip() != "":
                return float(str(value).replace(",", "").strip())
        except Exception:
            pass
    text = json.dumps(props, ensure_ascii=False).lower()
    import re
    matches = re.findall(r"(?:depth|water|น้ำ)[^0-9]{0,20}(\d+(?:\.\d+)?)\s*(?:cm|เซนติเมตร)?", text)
    if matches:
        try:
            return float(matches[0])
        except Exception:
            pass
    return None


def flood_status_from_properties(props):
    depth = normalize_depth(props)
    text = json.dumps(props or {}, ensure_ascii=False).lower()
    if depth is not None:
        if depth >= 70:
            return "severe", depth
        if depth >= 20:
            return "moderate", depth
        return "low", depth
    if any(x in text for x in ["not passable", "deep", "70+", "70 cm", "danger"]):
        return "severe", None
    if any(x in text for x in ["risky", "unconfirmed", "small car", "passable"]):
        return "moderate", None
    if any(x in text for x in ["receded", "water gone"]):
        return "low", 0
    return "unknown", None


def cached_get(url, filename, ttl=60):
    path = os.path.join(CACHE_DIR, filename)
    if os.path.exists(path) and time.time() - os.path.getmtime(path) < ttl:
        with open(path, "rb") as f:
            return f.read(), None
    r = requests.get(url, timeout=45, headers={"User-Agent": "FloodSiteBaseDashboard/2.0"})
    r.raise_for_status()
    content = r.content
    with open(path, "wb") as f:
        f.write(content)
    return content, r.headers.get("Content-Type")


def load_flood_index(force=False):
    """Build a spatial index from FloodBoard roads. Site coordinates are WGS84.
    A small degree buffer is used for proximity classification; this is intended
    as an operational indicator, not a hydrological inundation model.
    """
    global _flood_index, _flood_index_loaded_at, _flood_features
    if Point is None or STRtree is None:
        return None, []
    if _flood_index is not None and not force and time.time() - _flood_index_loaded_at < 60:
        return _flood_index, _flood_features
    try:
        content, _ = cached_get(FLOODBOARD_ROADS, "roads.geojson", ttl=60)
        gj = json.loads(content.decode("utf-8"))
        geoms, feats = [], []
        for f in gj.get("features", []):
            try:
                g = shape(f.get("geometry"))
                if not g.is_empty:
                    geoms.append(g)
                    feats.append(f)
            except Exception:
                continue
        _flood_features = feats
        _flood_index = STRtree(geoms) if geoms else None
        _flood_index_loaded_at = time.time()
        return _flood_index, _flood_features
    except Exception:
        return None, []


def classify_site(lat, lon, index, features):
    if index is None or not features:
        return "unknown", None, None
    p = Point(float(lon), float(lat))
    # Roughly 100 m at Thailand latitudes; used only to relate a Site to a reported road.
    candidate_idx = index.query(p.buffer(0.0012))
    best_status, best_depth, best_distance = "unknown", None, 999.0
    for idx in candidate_idx:
        try:
            geom = index.geometries[idx]
            distance = p.distance(geom)
            if distance > 0.0015:
                continue
            props = features[idx].get("properties", {}) or {}
            status, depth = flood_status_from_properties(props)
            priority = {"severe": 4, "moderate": 3, "low": 2, "unknown": 1}.get(status, 0)
            current_priority = {"severe": 4, "moderate": 3, "low": 2, "unknown": 1}.get(best_status, 0)
            if priority > current_priority or (priority == current_priority and distance < best_distance):
                best_status, best_depth, best_distance = status, depth, distance
        except Exception:
            continue
    if best_status == "unknown":
        return "safe", None, None
    # For display/API, treat an explicit low-water/receded report as low risk.
    return best_status, best_depth, best_distance * 111000.0



def normalize_col_name(v):
    return re.sub(r"[^a-z0-9ก-๙]+", "", clean(v).lower())


def pick_column(columns, candidates, contains=None):
    normalized = {normalize_col_name(c): c for c in columns}
    for c in candidates:
        if normalize_col_name(c) in normalized:
            return normalized[normalize_col_name(c)]
    if contains:
        for c in columns:
            n = normalize_col_name(c)
            if all(x in n for x in contains):
                return c
    return None


def parse_coords(text):
    t = clean(text)
    patterns = [
        r"(?:lat(?:itude)?)[^0-9-]{0,12}(-?\d+(?:\.\d+)?)\D{1,12}(?:lon(?:gitude)?)[^0-9-]{0,12}(-?\d+(?:\.\d+)?)",
        r"(-?\d{1,2}\.\d{3,})\s*[,;/| ]\s*(-?\d{2,3}\.\d{3,})",
    ]
    for pat in patterns:
        m=re.search(pat,t,re.I)
        if not m: continue
        try:
            lat,lon=float(m.group(1)),float(m.group(2))
            if 5 <= lat <= 21 and 97 <= lon <= 106:
                return lat,lon,"title"
        except Exception: pass
    return None,None,None


def parse_job_fields(title, explicit_job_id="", explicit_site_code=""):
    t=clean(title)
    job_id=clean(explicit_job_id)
    site_code=clean(explicit_site_code)
    if not job_id:
        patterns=[
            r"\bjob\s*(?:id|no|number|#)?\s*[:=\-#]?\s*([A-Za-z0-9][A-Za-z0-9._/-]{2,})",
            r"\bJOB[-_/#:]?([A-Za-z0-9][A-Za-z0-9._/-]{2,})",
        ]
        for pat in patterns:
            m=re.search(pat,t,re.I)
            if m:
                job_id=m.group(1).strip(' ,;|')
                break
    if not site_code:
        patterns=[
            r"\bsite\s*(?:code|id|no)?\s*[:=\-#]?\s*([A-Za-z0-9][A-Za-z0-9._/-]{2,})",
            r"\b(?:SITE|ST)[-_/#:]([A-Za-z0-9][A-Za-z0-9._/-]{2,})",
        ]
        for pat in patterns:
            m=re.search(pat,t,re.I)
            if m:
                site_code=m.group(1).strip(' ,;|')
                break
    lat,lon,coord_source=parse_coords(t)
    return job_id,site_code,lat,lon,coord_source


def enrich_job_flood(row, index_obj=None, features=None):
    """Attach current FloodBoard status/depth to a Job coordinate."""
    lat, lon = row.get("latitude"), row.get("longitude")
    if lat is None or lon is None:
        row["flood_status"] = "unknown"
        row["water_depth_cm"] = None
        row["flood_distance_m"] = None
        return row
    if index_obj is None and features is None:
        index_obj, features = load_flood_index()
    status, depth, distance = classify_site(lat, lon, index_obj, features)
    row["flood_status"] = status
    row["water_depth_cm"] = depth
    row["flood_distance_m"] = distance
    return row


def load_job_monitor():
    if not os.path.exists(JOB_CACHE): return {"filename":"","uploaded_at":"","title_column":"","rows":[]}
    try:
        with open(JOB_CACHE,'r',encoding='utf-8') as f: return json.load(f)
    except Exception:
        return {"filename":"","uploaded_at":"","title_column":"","rows":[]}


def save_job_monitor(data):
    with open(JOB_CACHE,'w',encoding='utf-8') as f:
        json.dump(data,f,ensure_ascii=False,indent=2)


def enrich_job_site(row):
    code=clean(row.get("site_code"))
    if not code: return row
    conn=db()
    try:
        s=conn.execute("SELECT site_code,location_name,tumbol,amphur,province,latitude,longitude FROM sites WHERE site_code=? LIMIT 1",(code,)).fetchone()
        if s:
            sd=dict(s)
            row["site_match"]=True
            row["site_location_name"]=sd["location_name"]
            row["site_province"]=sd["province"]
            row["site_amphur"]=sd["amphur"]
            if row.get("latitude") is None and row.get("longitude") is None:
                row["latitude"],row["longitude"]=sd["latitude"],sd["longitude"]
                row["coord_source"]="site_base"
        else: row["site_match"]=False
    finally: conn.close()
    return row


def get_sites(filters=None, include_flood=False, limit=None):
    filters = filters or {}
    province = clean(filters.get("province"))
    amphur = clean(filters.get("amphur"))
    q = clean(filters.get("q"))
    south, west, north, east = [filters.get(k) for k in ("south", "west", "north", "east")]
    conn = db()
    try:
        where = ["latitude IS NOT NULL", "longitude IS NOT NULL", "province IN (?,?,?,?)"]
        params = list(ALLOWED_PROVINCES)
        if province:
            where.append("province = ?"); params.append(province)
        if amphur:
            where.append("amphur = ?"); params.append(amphur)
        if q:
            where.append("(site_code LIKE ? OR location_name LIKE ? OR tumbol LIKE ? OR amphur LIKE ? OR province LIKE ?)")
            like = f"%{q}%"; params.extend([like] * 5)
        if None not in (south, west, north, east):
            where += ["latitude BETWEEN ? AND ?", "longitude BETWEEN ? AND ?"]
            params += [float(south), float(north), float(west), float(east)]
        sql = f"SELECT site_code, location_name, tumbol, amphur, province, latitude, longitude FROM sites WHERE {' AND '.join(where)} ORDER BY site_code"
        if limit is not None:
            sql += " LIMIT ?"; params.append(int(limit))
        return [row_to_dict(r) for r in conn.execute(sql, params).fetchall()]
    finally:
        conn.close()


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/filters")
def filters():
    conn = db()
    try:
        provinces = [p for p in ALLOWED_PROVINCES if conn.execute("SELECT 1 FROM sites WHERE province=? LIMIT 1", (p,)).fetchone()]
        province = clean(request.args.get("province"))
        if province:
            amphurs = [r[0] for r in conn.execute("SELECT DISTINCT amphur FROM sites WHERE province=? AND amphur IS NOT NULL AND amphur<>'' ORDER BY amphur", (province,)).fetchall()]
        else:
            amphurs = [r[0] for r in conn.execute("SELECT DISTINCT amphur FROM sites WHERE province IN (?,?,?,?) AND amphur IS NOT NULL AND amphur<>'' ORDER BY amphur", tuple(ALLOWED_PROVINCES)).fetchall()]
        return jsonify({"provinces": provinces, "amphurs": amphurs})
    finally:
        conn.close()


@app.get("/api/sites")
def sites():
    limit = min(max(request.args.get("limit", 5000, type=int), 1), 10000)
    data = get_sites(request.args, include_flood=False, limit=limit)
    if request.args.get("with_flood", "0") == "1":
        index_obj, features = load_flood_index()
        for s in data:
            status, depth, distance = classify_site(s["latitude"], s["longitude"], index_obj, features)
            s.update(flood_status=status, water_depth_cm=depth, flood_distance_m=distance)
    return jsonify({"count": len(data), "limit": limit, "sites": data})


@app.get("/api/site/<site_code>")
def site(site_code):
    conn = db()
    try:
        row = conn.execute("SELECT * FROM sites WHERE site_code = ? AND province IN (?,?,?,?) LIMIT 1", (site_code, *ALLOWED_PROVINCES)).fetchone()
        if not row:
            return jsonify({"error": "Site not found"}), 404
        result = row_to_dict(row)
        index_obj, features = load_flood_index()
        status, depth, distance = classify_site(result["latitude"], result["longitude"], index_obj, features)
        result.update(flood_status=status, water_depth_cm=depth, flood_distance_m=distance)
        return jsonify(result)
    finally:
        conn.close()


@app.get("/api/site-count")
def site_count():
    conn = db()
    try:
        total = conn.execute("SELECT COUNT(*) AS n FROM sites WHERE province IN (?,?,?,?)", tuple(ALLOWED_PROVINCES)).fetchone()["n"]
        provinces = conn.execute("SELECT COUNT(DISTINCT province) AS n FROM sites WHERE province IN (?,?,?,?)", tuple(ALLOWED_PROVINCES)).fetchone()["n"]
        return jsonify({"total": total, "provinces": provinces})
    finally:
        conn.close()


@app.get("/api/dashboard-summary")
def dashboard_summary():
    sites_data = get_sites(request.args, limit=None)
    index_obj, features = load_flood_index()
    counts = {"total": len(sites_data), "flooded": 0, "severe": 0, "moderate": 0, "low": 0, "safe": 0, "unknown": 0}
    for s in sites_data:
        status, depth, distance = classify_site(s["latitude"], s["longitude"], index_obj, features)
        counts[status] = counts.get(status, 0) + 1
    counts["flooded"] = counts.get("severe", 0) + counts.get("moderate", 0) + counts.get("low", 0)
    counts["matched"] = counts["total"]
    return jsonify(counts)



@app.post("/api/job-monitor/upload")
def job_monitor_upload():
    f=request.files.get("file")
    if not f or not f.filename:
        return jsonify({"error":"กรุณาเลือกไฟล์ Excel Job Monitor"}),400
    if not f.filename.lower().endswith((".xlsx",".xlsm",".xltx",".xltm")):
        return jsonify({"error":"รองรับไฟล์ Excel .xlsx/.xlsm/.xltx/.xltm"}),400
    try:
        import pandas as pd
        df=pd.read_excel(f, sheet_name=0, dtype=object)
        df.columns=[clean(c) for c in df.columns]
        cols=list(df.columns)
        title_col=pick_column(cols,["Title Job","Job Title","Title","Job Name","JobName"],contains=["title"])
        job_col=pick_column(cols,["Job ID","Job_ID","JobID","Job No","Job Number","ID"])
        site_col=pick_column(cols,["Site Code","Site_Code","SiteCode","SITE_CODE","Site ID"])
        if not title_col:
            return jsonify({"error":"หา column Title Job ไม่พบ","columns":cols}),400
        conn=db()
        try:
            site_map={r["site_code"]:dict(r) for r in conn.execute("SELECT site_code,location_name,tumbol,amphur,province,latitude,longitude FROM sites WHERE province IN (?,?,?,?)",tuple(ALLOWED_PROVINCES)).fetchall()}
        finally:
            conn.close()
        rows=[]
        flood_index, flood_features = load_flood_index()
        for i,rec in df.iterrows():
            title=clean(rec.get(title_col,''))
            job_id,site_code,lat,lon,coord_source=parse_job_fields(title, rec.get(job_col,'') if job_col else '', rec.get(site_col,'') if site_col else '')
            row={"row_number":int(i)+2,"title":title,"job_id":job_id,"site_code":site_code,"latitude":lat,"longitude":lon,"coord_source":coord_source or ""}
            sd=site_map.get(site_code) if site_code else None
            if sd:
                row["site_match"]=True
                row["site_location_name"]=sd["location_name"]
                row["site_province"]=sd["province"]
                row["site_amphur"]=sd["amphur"]
                if row["latitude"] is None or row["longitude"] is None:
                    row["latitude"],row["longitude"]=sd["latitude"],sd["longitude"]
                    row["coord_source"]="site_base"
            else:
                row["site_match"]=False
            row["parse_status"]="ok" if (row["job_id"] or row["site_code"] or row["latitude"] is not None) else "unparsed"
            enrich_job_flood(row, flood_index, flood_features)
            rows.append(row)
        data={"filename":f.filename,"uploaded_at":time.strftime("%Y-%m-%d %H:%M:%S"),"title_column":title_col,"job_column":job_col or "","site_column":site_col or "","rows":rows}
        save_job_monitor(data)
        return jsonify({"filename":f.filename,"title_column":title_col,"job_column":job_col,"site_column":site_col,"count":len(rows),"matched":sum(1 for x in rows if x.get('site_match')),"with_coords":sum(1 for x in rows if x.get('latitude') is not None and x.get('longitude') is not None),"rows":rows[:1000]})
    except Exception as e:
        return jsonify({"error":"อ่านไฟล์ Excel ไม่สำเร็จ","detail":str(e)}),400


@app.get("/api/job-monitor")
def job_monitor():
    data=load_job_monitor()
    rows=data.get("rows",[])
    q=clean(request.args.get("q"))
    if q:
        ql=q.lower(); rows=[r for r in rows if ql in clean(r.get("title")).lower() or ql in clean(r.get("job_id")).lower() or ql in clean(r.get("site_code")).lower()]
    flood_index, flood_features = load_flood_index()
    for r in rows:
        if "flood_status" not in r:
            enrich_job_flood(r, flood_index, flood_features)
    return jsonify({"filename":data.get("filename",""),"uploaded_at":data.get("uploaded_at",""),"title_column":data.get("title_column",""),"count":len(rows),"rows":rows[:5000]})


@app.get("/api/job-monitor/export.xlsx")
def job_monitor_export():
    data=load_job_monitor(); rows=data.get("rows",[])
    wb=Workbook(); ws=wb.active; ws.title="Job Monitor Analysis"
    headers=["Row","Job ID","Site Code","Title Job","Latitude","Longitude","Coordinate Source","Site Match","Site Location","Amphur","Province","Flood Status","Water Depth (cm)","Distance to Flooded Road (m)","Parse Status"]
    ws.append(headers)
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1769AA"); c.alignment=Alignment(horizontal="center")
    for r in rows:
        ws.append([r.get("row_number"),r.get("job_id"),r.get("site_code"),r.get("title"),r.get("latitude"),r.get("longitude"),r.get("coord_source"),"YES" if r.get("site_match") else "NO",r.get("site_location_name"),r.get("site_amphur"),r.get("site_province"),r.get("flood_status"),r.get("water_depth_cm"),r.get("flood_distance_m"),r.get("parse_status")])
    widths=[8,20,20,70,14,14,18,12,35,22,24,16,18,28,14]
    for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width=w
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    buf=BytesIO(); wb.save(buf); buf.seek(0)
    stamp=time.strftime("%Y%m%d_%H%M%S")
    return send_file(buf,as_attachment=True,download_name=f"job_monitor_analysis_{stamp}.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/export.xlsx")
def export_xlsx():
    rows = get_sites(request.args, limit=None)
    index_obj, features = load_flood_index()
    wb = Workbook()
    ws = wb.active
    ws.title = "Site Base Report"
    headers = ["Site Code", "Location Name", "Tambon", "Amphur", "Province", "Latitude", "Longitude", "Flood Status", "Water Depth (cm)", "Distance to Flooded Road (m)"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1769AA")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for s in rows:
        status, depth, distance = classify_site(s["latitude"], s["longitude"], index_obj, features)
        ws.append([
            s["site_code"], s["location_name"], s["tumbol"], s["amphur"], s["province"],
            s["latitude"], s["longitude"], status, depth, distance
        ])
    widths = [16, 48, 22, 22, 24, 13, 13, 16, 18, 28]
    for i, w in enumerate(widths, 1): ws.column_dimensions[chr(64+i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = BytesIO(); wb.save(buf); buf.seek(0)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    return send_file(buf, as_attachment=True, download_name=f"site_base_flood_report_{stamp}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/floodboard/roads")
def floodboard_roads():
    try:
        content, _ = cached_get(FLOODBOARD_ROADS, "roads.geojson", ttl=60)
        return Response(content, content_type="application/geo+json")
    except Exception as e:
        return jsonify({"error": "Cannot load FloodBoard roads data", "detail": str(e)}), 502


@app.get("/api/floodboard/reports")
def floodboard_reports():
    try:
        content, _ = cached_get(FLOODBOARD_REPORTS, "reports.csv", ttl=60)
        return Response(content, content_type="text/csv; charset=utf-8")
    except Exception as e:
        return jsonify({"error": "Cannot load FloodBoard reports data", "detail": str(e)}), 502


@app.get("/api/floodboard/stats")
def floodboard_stats():
    try:
        content, _ = cached_get(FLOODBOARD_STATS, "stats.json", ttl=60)
        return Response(content, content_type="application/json; charset=utf-8")
    except Exception as e:
        return jsonify({"error": "Cannot load FloodBoard stats", "detail": str(e)}), 502


@app.get("/api/floodboard/feed")
def floodboard_feed():
    try:
        content, _ = cached_get(FLOODBOARD_FEED, "feed.json", ttl=60)
        return Response(content, content_type="application/json; charset=utf-8")
    except Exception as e:
        return jsonify({"error": "Cannot load FloodBoard feed", "detail": str(e)}), 502


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=True)
